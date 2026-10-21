import sys
import zipfile

import pytest
from conftest import criar_pdf_com_texto
from fpdf import FPDF
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

from unepdf.core import pdf_office
from unepdf.core.erros import ErroUnePDF
from unepdf.core.progresso import Cancelamento
from unepdf.utils.programas import office_instalado

LINHAS = [("Nota", "Emitente", "Valor"), ("1204", "EMPRESA EXEMPLO LTDA", "1.234,56"), ("0007", "OUTRA EMPRESA SA", "89,90")]


def relatorio(caminho, paginas=1, bordas=1, titulo="Relatorio de notas do mes"):
    """PDF com uma tabela por página; sem bordas, as colunas só ficam alinhadas."""
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for _ in range(paginas):
        pdf.add_page()
        pdf.cell(0, 10, titulo, new_x="LMARGIN", new_y="NEXT")
        for linha in LINHAS:
            for texto, largura in zip(linha, (30, 90, 40)):
                pdf.cell(largura, 8, texto, border=bordas)
            pdf.ln()
    caminho.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(caminho))
    return caminho


def _linhas(arquivo, aba=0):
    return [linha for linha in load_workbook(arquivo).worksheets[aba].iter_rows(values_only=True)]


@pytest.mark.parametrize("texto, valor", [
    ("1.234,56", 1234.56), ("R$ 89,90", 89.9), ("-12,5", -12.5), ("1.234", 1234), ("42", 42), ("0", 0),
    ("0007", "0007"), ("11.222.333/0001-81", "11.222.333/0001-81"), ("15/01/2026", "15/01/2026"),
    ("35260111222333000181550010000012341000012345", "35260111222333000181550010000012341000012345"),
    ("  EMPRESA   EXEMPLO ", "EMPRESA EXEMPLO"), ("", None), (None, None),
])
def test_valor_da_celula(texto, valor):
    assert pdf_office.valor_da_celula(texto) == valor


# PDF para Excel

def test_tabela_com_linhas_vira_planilha(tmp_path):
    origem = relatorio(tmp_path / "notas.pdf")
    resultado = pdf_office.pdf_para_excel([origem], tmp_path / "UnePDF")

    assert resultado.gerados == [tmp_path / "UnePDF" / "notas.xlsx"] and not resultado.ignorados
    assert load_workbook(resultado.gerados[0]).sheetnames == ["Página 1"]
    assert _linhas(resultado.gerados[0]) == [
        ("Nota", "Emitente", "Valor"), (1204, "EMPRESA EXEMPLO LTDA", 1234.56), ("0007", "OUTRA EMPRESA SA", 89.9),
    ]


def test_tabela_sem_linhas_pelo_alinhamento(tmp_path):
    origem = relatorio(tmp_path / "notas.pdf", bordas=0)
    linhas = _linhas(pdf_office.pdf_para_excel([origem], tmp_path / "UnePDF").gerados[0])
    # O título atravessava a divisa das colunas: as partes da palavra cortada voltam a se juntar
    assert linhas[0][0] == "Relatorio de notas do mes"
    assert (1204, "EMPRESA EXEMPLO LTDA", 1234.56) in linhas
    assert all(any(celula is not None for celula in linha) for linha in linhas)


def test_uma_aba_por_pagina_ou_todas_juntas(tmp_path):
    origem = relatorio(tmp_path / "notas.pdf", paginas=2)
    por_pagina = pdf_office.pdf_para_excel([origem], tmp_path / "a").gerados[0]
    assert load_workbook(por_pagina).sheetnames == ["Página 1", "Página 2"]

    juntas = pdf_office.pdf_para_excel([origem], tmp_path / "b", uma_aba=True).gerados[0]
    assert load_workbook(juntas).sheetnames == ["Tabelas"]
    assert _linhas(juntas).count(("Nota", "Emitente", "Valor")) == 2


def test_pagina_sem_tabela_fica_de_fora(tmp_path):
    origem = relatorio(tmp_path / "notas.pdf")
    texto = criar_pdf_com_texto(tmp_path / "texto.pdf", [""])
    escritor = PdfWriter()
    for pagina in [*PdfReader(origem).pages, *PdfReader(texto).pages]:
        escritor.add_page(pagina)
    escritor.write(tmp_path / "misto.pdf")

    resultado = pdf_office.pdf_para_excel([tmp_path / "misto.pdf"], tmp_path / "UnePDF")
    assert resultado.paginas_sem_tabela == 1
    assert load_workbook(resultado.gerados[0]).sheetnames == ["Página 1"]


def test_pdfs_que_nao_viram_planilha(tmp_path, criar_pdf):
    escaneado = criar_pdf_com_texto(tmp_path / "escaneado.pdf", [""])
    protegido = criar_pdf(tmp_path / "protegido.pdf", senha="1234")
    bom = relatorio(tmp_path / "bom.pdf")

    resultado = pdf_office.pdf_para_excel([escaneado, protegido, bom], tmp_path / "UnePDF")
    assert [g.name for g in resultado.gerados] == ["bom.xlsx"]
    assert {i.caminho.name: i.motivo for i in resultado.ignorados} == {
        "escaneado.pdf": pdf_office.SEM_TEXTO, "protegido.pdf": "protegido por senha",
    }

    with pytest.raises(ErroUnePDF) as erro:
        pdf_office.pdf_para_excel([escaneado], tmp_path / "UnePDF")
    assert erro.value.ignorados[0].motivo == pdf_office.SEM_TEXTO


def test_cancelar_a_conversao_para_excel(tmp_path):
    origens = [relatorio(tmp_path / f"{n}.pdf") for n in (1, 2)]
    cancelamento = Cancelamento()
    cancelamento.cancelar()
    resultado = pdf_office.pdf_para_excel(origens, tmp_path / "UnePDF", cancelamento=cancelamento)
    assert resultado.cancelado and resultado.pendentes == 2 and not resultado.gerados


# PDF para Word

class _RegistroFalso:
    HKEY_CURRENT_USER = "HKCU"
    REG_DWORD = 4

    def __init__(self, valores=None):
        self.valores = dict(valores or {})

    class _Chave:
        def __init__(self, nome):
            self.nome = nome

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    def CreateKey(self, raiz, nome):
        return self._Chave(nome)

    def QueryValueEx(self, chave, valor):
        if (chave.nome, valor) not in self.valores:
            raise OSError("não existe")
        return self.valores[chave.nome, valor]

    def SetValueEx(self, chave, valor, reservado, tipo, dado):
        self.valores[chave.nome, valor] = (dado, tipo)

    def DeleteValue(self, chave, valor):
        del self.valores[chave.nome, valor]


CHAVE_WORD = r"Software\Microsoft\Office\16.0\Word\Options"


@pytest.mark.parametrize("antes", [{}, {(CHAVE_WORD, "DisableConvertPdfWarning"): (0, 4)}])
def test_aviso_de_conversao_volta_ao_valor_original(monkeypatch, antes):
    registro = _RegistroFalso(antes)
    monkeypatch.setitem(sys.modules, "winreg", registro)
    with pdf_office._sem_aviso_de_conversao("16.0"):
        assert registro.valores[CHAVE_WORD, "DisableConvertPdfWarning"] == (1, 4)
    assert registro.valores == antes


def test_pdf_escaneado_nem_chega_ao_word(tmp_path, monkeypatch):
    abertos = []
    monkeypatch.setattr(pdf_office.office._InstanciaOffice, "abrir", lambda self: abertos.append(self))
    escaneado = criar_pdf_com_texto(tmp_path / "escaneado.pdf", [""])
    with pytest.raises(ErroUnePDF) as erro:
        pdf_office.pdf_para_word([escaneado], tmp_path / "UnePDF")
    assert erro.value.ignorados[0].motivo == pdf_office.SEM_TEXTO
    assert abertos == []


def test_word_recebe_uma_copia_sem_a_senha_de_permissao(tmp_path, criar_pdf):
    texto = criar_pdf_com_texto(tmp_path / "contrato.pdf", ["CONTRATO DE PRESTACAO DE SERVICOS ENTRE AS PARTES"])
    escritor = PdfWriter(clone_from=PdfReader(texto))
    escritor.encrypt(user_password="", owner_password="dono", algorithm="AES-256")
    escritor.write(tmp_path / "restrito.pdf")

    copia = pdf_office._copia_legivel(tmp_path / "restrito.pdf", tmp_path)
    assert copia.name.startswith("unepdf-") and copia.parent == tmp_path
    assert not PdfReader(copia).is_encrypted


sem_word = pytest.mark.skipif(not office_instalado("Word.Application"), reason="Microsoft Word não instalado")


@sem_word
def test_pdf_vira_documento_do_word(tmp_path):
    origem = criar_pdf_com_texto(tmp_path / "contrato.pdf", [
        "CONTRATO DE PRESTACAO DE SERVICOS ENTRE AS PARTES", "Clausula segunda: o pagamento sera mensal.",
    ])
    resultado = pdf_office.pdf_para_word([origem], tmp_path / "UnePDF")
    assert resultado.gerados == [tmp_path / "UnePDF" / "contrato.docx"]
    documento = zipfile.ZipFile(resultado.gerados[0]).read("word/document.xml").decode("utf-8")
    assert "CONTRATO DE PRESTACAO" in documento and "pagamento sera mensal" in documento
