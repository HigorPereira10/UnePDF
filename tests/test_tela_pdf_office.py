import pytest
from openpyxl import load_workbook
from test_pdf_office import relatorio

from unepdf.core import pdf_office
from unepdf.ui.contexto import ContextoApp
from unepdf.ui.janela_principal import JanelaPrincipal


@pytest.fixture
def janela(qtbot, tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    janela = JanelaPrincipal(ContextoApp())
    qtbot.addWidget(janela)
    janela.show()
    return janela


def _esperar(qtbot, tela):
    qtbot.waitUntil(lambda: not tela.executor.em_execucao, timeout=20_000)


def test_pdf_para_excel_pela_tela(qtbot, janela, tmp_path):
    janela.abrir("pdf_excel")
    tela = janela.telas["pdf_excel"]
    relatorio(tmp_path / "notas" / "janeiro.pdf", paginas=2)

    tela.receber_caminhos([tmp_path / "notas"])
    _esperar(qtbot, tela)
    assert tela.estados.currentWidget() is tela.painel_selecao
    assert tela.lista.item(0).text() == "janeiro.pdf  >  janeiro.xlsx"
    assert tela.destino_selecao.text() == str(tmp_path / "notas" / "UnePDF")

    tela.check_uma_aba.setChecked(True)
    assert tela.contexto.config.excel_uma_aba is True
    tela.botao_converter.click()
    _esperar(qtbot, tela)
    assert tela.estados.currentWidget() is tela.resultado
    assert load_workbook(tmp_path / "notas" / "UnePDF" / "janeiro.xlsx").sheetnames == ["Tabelas"]


def test_pdf_para_word_sem_word_fica_desativado(qtbot, janela, monkeypatch):
    monkeypatch.setattr(pdf_office, "office_instalado", lambda prog_id: False)
    janela.abrir("pdf_word")
    tela = janela.telas["pdf_word"]
    assert not tela.area.botao_arquivos.isEnabled()
    assert "Microsoft Word" in tela.area.aviso.text()


def test_pdf_para_word_pela_tela(qtbot, janela, tmp_path, monkeypatch, criar_pdf):
    monkeypatch.setattr(pdf_office, "office_instalado", lambda prog_id: True)
    chamadas = []

    def converter(arquivos, pasta, progresso, cancelamento):
        chamadas.append((arquivos, pasta))
        return pdf_office.ResultadoConversao(pasta, gerados=[pasta / "contrato.docx"])

    monkeypatch.setattr(pdf_office, "pdf_para_word", converter)
    janela.abrir("pdf_word")
    tela = janela.telas["pdf_word"]
    criar_pdf(tmp_path / "contrato.pdf")

    tela.receber_caminhos([tmp_path / "contrato.pdf"])
    _esperar(qtbot, tela)
    tela.botao_converter.click()
    _esperar(qtbot, tela)
    assert chamadas == [([tmp_path / "contrato.pdf"], tmp_path / "UnePDF")]
    assert tela.estados.currentWidget() is tela.resultado
